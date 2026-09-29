"""Workbook parsing and amount matching for a drop folder of invoices."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

TOLERANCE = Decimal("0.01")
_MONEY = Decimal("0.01")
_MAX_INVOICE_PARTS = 6

_FILENAME = re.compile(
    r"^(?P<boat>.+)-(?P<day>\d{1,2})-(?P<mon>[A-Za-z]{3})-(?P<year>\d{4})-\d{2}_\d{2}$"
)
_MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
_DATE_FORMATS = ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")

REQUIRED_COLUMNS = (
    "SUPPLIER NAME",
    "INVOICE NUMBER",
    "DESCRIPTION",
    "AMOUNT ORIGINAL",
    "CURRENCY ORIGIN",
    "CREATE DATE",
    "INVOICE/QUESTION DATE",
    "WORK/ORDER STATUS",
    "BOAT CATEGORY",
    "BOAT SUBCATEGORY",
)


@dataclass(frozen=True)
class PaymentRow:
    row_number: int
    supplier_name: str
    invoice_number: str
    description: str
    amount_original: Decimal
    currency: str
    create_date: date | None
    invoice_date: date | None
    work_order_status: str
    boat_category: str
    boat_subcategory: str


@dataclass(frozen=True)
class ExtractedInvoice:
    source_name: str
    total_amount: Decimal
    invoice_number: str | None = None
    net_amount: Decimal | None = None
    vat_amount: Decimal | None = None
    supplier_name: str | None = None
    invoice_date: str | None = None
    currency: str | None = None
    description: str | None = None
    drive_url: str | None = None


@dataclass(frozen=True)
class MatchedPayment:
    payment: PaymentRow
    invoices: tuple[ExtractedInvoice, ...]

    @property
    def matched(self) -> bool:
        return bool(self.invoices)


def parse_amount(value: object) -> Decimal:
    """Parse a workbook or model amount and round it to cents."""
    if isinstance(value, Decimal):
        amount = value
    elif isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid amount {value!r}")
    elif isinstance(value, int):
        amount = Decimal(value)
    elif isinstance(value, float):
        amount = Decimal(str(value))
    else:
        text = _normalize_amount_text(str(value))
        if not text:
            raise ValueError("Invalid amount ''")
        try:
            amount = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(f"Invalid amount {value!r}") from exc
    return amount.quantize(_MONEY, rounding=ROUND_HALF_UP)


def format_money(amount: Decimal) -> str:
    quantized = amount.quantize(_MONEY, rounding=ROUND_HALF_UP)
    return f"{quantized:.2f}"


def parse_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date {text!r}")


def format_date(value: date | None) -> str:
    if value is None:
        return ""
    return value.isoformat()


def display_boat_name(token: str) -> str:
    """Turn a filename boat token into a Drive folder name."""
    name = token.replace("_s_", "'s ").replace("_", " ")
    return " ".join(name.split())


def parse_workbook_filename(stem: str) -> tuple[str, date] | None:
    """Return the boat folder name and archive date encoded in a workbook stem."""
    match = _FILENAME.fullmatch(stem)
    if match is None:
        return None
    month = _MONTHS.get(match.group("mon").lower())
    if month is None:
        raise ValueError(f"Unrecognized month {match.group('mon')!r} in {stem}")
    try:
        archive_date = date(int(match.group("year")), month, int(match.group("day")))
    except ValueError as exc:
        raise ValueError(f"Invalid date in workbook name {stem}") from exc
    return display_boat_name(match.group("boat")), archive_date


def resolve_archive_target(
    workbook_name: str,
    boat_override: str | None,
    today: date,
) -> tuple[str, date]:
    parsed = parse_workbook_filename(Path(workbook_name).stem)
    if parsed is not None:
        return parsed
    override = (boat_override or "").strip()
    if override:
        return override, today
    raise ValueError(
        f"Could not read a boat and date from {workbook_name}. "
        "Name it like Boat_Name-18-Sep-2026-06_38.xlsx or pass --boat."
    )


def discover_inputs(directory: Path) -> tuple[Path, list[Path]]:
    """Find the single workbook and every PDF in a drop folder."""
    if not directory.is_dir():
        raise ValueError(f"Input directory does not exist: {directory}")
    files = [
        path
        for path in directory.iterdir()
        if path.is_file() and not path.name.startswith(".") and not path.name.startswith("~$")
    ]
    workbooks = sorted(path for path in files if path.suffix.lower() == ".xlsx")
    pdfs = sorted(path for path in files if path.suffix.lower() == ".pdf")
    if len(workbooks) != 1:
        found = ", ".join(path.name for path in workbooks) or "none"
        raise ValueError(
            f"Expected exactly one .xlsx in {directory}, found {len(workbooks)}: {found}"
        )
    return workbooks[0], pdfs


def load_payments(path: Path) -> list[PaymentRow]:
    """Read outstanding-payment rows from the first worksheet."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError(f"{path.name} is empty") from exc
        indexes = _header_indexes(header)
        payments: list[PaymentRow] = []
        for row_number, values in enumerate(rows, start=2):
            if values is None or all(_blank(cell) for cell in values):
                continue
            payments.append(_payment_from_row(row_number, values, indexes))
        return payments
    finally:
        workbook.close()


def split_invoice_numbers(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[/&,]", value or "") if part.strip()]


def match_payments(
    payments: list[PaymentRow],
    invoices: list[ExtractedInvoice],
) -> list[MatchedPayment]:
    """Match each payment to unused invoices by total, then by summed totals."""
    unused = {id(invoice): invoice for invoice in invoices}
    matched: list[MatchedPayment] = []
    for payment in payments:
        available = list(unused.values())
        chosen = _match_one(payment, available) or _match_many(payment, available)
        if not chosen:
            matched.append(MatchedPayment(payment, ()))
            continue
        for invoice in chosen:
            unused.pop(id(invoice), None)
        ordered = tuple(sorted(chosen, key=lambda invoice: invoice.source_name))
        matched.append(MatchedPayment(payment, ordered))
    return matched


def _normalize_amount_text(text: str) -> str:
    cleaned = text.strip().replace(" ", "").replace("\u00a0", "")
    cleaned = cleaned.replace("€", "").replace("$", "")
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    return cleaned


def _blank(value: object) -> bool:
    return value is None or str(value).strip() == ""


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_indexes(header: tuple[object, ...] | None) -> dict[str, int]:
    if not header:
        raise ValueError("Workbook is missing its header row")
    indexes: dict[str, int] = {}
    for index, cell in enumerate(header):
        name = _text(cell)
        if not name:
            continue
        if name in indexes:
            raise ValueError(f"Duplicate column {name}")
        indexes[name] = index
    missing = [name for name in REQUIRED_COLUMNS if name not in indexes]
    if missing:
        raise ValueError("Workbook is missing columns: " + ", ".join(missing))
    return indexes


def _cell(values: tuple[object, ...], indexes: dict[str, int], name: str) -> object:
    index = indexes[name]
    if index >= len(values):
        return None
    return values[index]


def _payment_from_row(
    row_number: int,
    values: tuple[object, ...],
    indexes: dict[str, int],
) -> PaymentRow:
    amount_raw = _cell(values, indexes, "AMOUNT ORIGINAL")
    if _blank(amount_raw):
        raise ValueError(f"Row {row_number} is missing AMOUNT ORIGINAL")
    try:
        amount = parse_amount(amount_raw)
    except ValueError as exc:
        raise ValueError(f"Row {row_number} AMOUNT ORIGINAL: {exc}") from exc
    try:
        create_date = parse_date(_cell(values, indexes, "CREATE DATE"))
    except ValueError as exc:
        raise ValueError(f"Row {row_number} CREATE DATE: {exc}") from exc
    try:
        invoice_date = parse_date(_cell(values, indexes, "INVOICE/QUESTION DATE"))
    except ValueError as exc:
        raise ValueError(f"Row {row_number} INVOICE/QUESTION DATE: {exc}") from exc
    return PaymentRow(
        row_number=row_number,
        supplier_name=_text(_cell(values, indexes, "SUPPLIER NAME")),
        invoice_number=_text(_cell(values, indexes, "INVOICE NUMBER")),
        description=_text(_cell(values, indexes, "DESCRIPTION")),
        amount_original=amount,
        currency=_text(_cell(values, indexes, "CURRENCY ORIGIN")),
        create_date=create_date,
        invoice_date=invoice_date,
        work_order_status=_text(_cell(values, indexes, "WORK/ORDER STATUS")),
        boat_category=_text(_cell(values, indexes, "BOAT CATEGORY")),
        boat_subcategory=_text(_cell(values, indexes, "BOAT SUBCATEGORY")),
    )


def _amounts_close(left: Decimal, right: Decimal) -> bool:
    return abs(left - right) <= TOLERANCE


def _digit_groups(text: str) -> list[str]:
    return re.findall(r"\d+", text or "")


def _same_number(left: str, right: str) -> bool:
    if left == right:
        return True
    left_norm = left.lstrip("0") or "0"
    right_norm = right.lstrip("0") or "0"
    return left_norm == right_norm


def overlap_score(invoice_number: str, *texts: str) -> int:
    """Count invoice-number parts that appear in the extracted number or filename."""
    parts = split_invoice_numbers(invoice_number)
    groups: list[str] = []
    blob_parts: list[str] = []
    for text in texts:
        if not text:
            continue
        blob_parts.append(text)
        groups.extend(_digit_groups(text))
    blob = " ".join(blob_parts).casefold()
    score = 0
    for part in parts:
        digits = _digit_groups(part)
        if digits:
            if any(_same_number(needle, group) for needle in digits for group in groups):
                score += 1
        elif part.casefold() in blob:
            score += 1
    return score


def _match_one(
    payment: PaymentRow,
    available: list[ExtractedInvoice],
) -> tuple[ExtractedInvoice, ...] | None:
    candidates = [
        invoice
        for invoice in available
        if _amounts_close(invoice.total_amount, payment.amount_original)
    ]
    if not candidates:
        return None
    if len(candidates) == 1:
        return (candidates[0],)
    ranked = sorted(
        candidates,
        key=lambda invoice: (
            -overlap_score(payment.invoice_number, invoice.invoice_number or "", invoice.source_name),
            invoice.source_name,
        ),
    )
    best = overlap_score(payment.invoice_number, ranked[0].invoice_number or "", ranked[0].source_name)
    second = overlap_score(
        payment.invoice_number,
        ranked[1].invoice_number or "",
        ranked[1].source_name,
    )
    if best <= 0 or best == second:
        return None
    return (ranked[0],)


def _match_many(
    payment: PaymentRow,
    available: list[ExtractedInvoice],
) -> tuple[ExtractedInvoice, ...] | None:
    parts = split_invoice_numbers(payment.invoice_number)
    # A shorter subset would not match the invoice list printed on the row.
    if len(parts) <= 1 or len(parts) > _MAX_INVOICE_PARTS:
        return None
    if len(available) < len(parts):
        return None
    best: list[tuple[ExtractedInvoice, ...]] = []
    best_score = -1
    for combo in combinations(available, len(parts)):
        total = sum((invoice.total_amount for invoice in combo), Decimal("0"))
        if not _amounts_close(total, payment.amount_original):
            continue
        numbers = " ".join(invoice.invoice_number or "" for invoice in combo)
        names = " ".join(invoice.source_name for invoice in combo)
        score = overlap_score(payment.invoice_number, numbers, names)
        ordered = tuple(sorted(combo, key=lambda invoice: invoice.source_name))
        if score > best_score:
            best_score = score
            best = [ordered]
        elif score == best_score:
            best.append(ordered)
    if len(best) != 1:
        return None
    return best[0]
