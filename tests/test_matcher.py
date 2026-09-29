"""Unit tests for workbook parsing, invoice matching, and CSV export."""

from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from services.csv_export import CSV_HEADERS, to_csv_fields, write_csv
from services.matcher import (
    ExtractedInvoice,
    PaymentRow,
    discover_inputs,
    load_payments,
    match_payments,
    parse_amount,
    parse_workbook_filename,
    resolve_archive_target,
)
from services.matcher import MatchedPayment
from services.pdf_extractor import InvoiceExtraction, retry_after_seconds, to_extracted_invoice


def _payment(**overrides: object) -> PaymentRow:
    values: dict[str, object] = {
        "row_number": 2,
        "supplier_name": "DERAMAR",
        "invoice_number": "100",
        "description": "service",
        "amount_original": Decimal("100.00"),
        "currency": "EUR",
        "create_date": date(2026, 9, 1),
        "invoice_date": date(2026, 9, 2),
        "work_order_status": "Completed",
        "boat_category": "Maintenance",
        "boat_subcategory": "Electrician",
    }
    values.update(overrides)
    return PaymentRow(**values)  # type: ignore[arg-type]


def _invoice(
    name: str,
    total: str,
    number: str | None = None,
    net: str | None = None,
    vat: str | None = None,
    url: str | None = None,
) -> ExtractedInvoice:
    return ExtractedInvoice(
        source_name=name,
        total_amount=Decimal(total),
        invoice_number=number,
        net_amount=None if net is None else Decimal(net),
        vat_amount=None if vat is None else Decimal(vat),
        drive_url=url,
    )


class FilenameTests(unittest.TestCase):
    def test_worlds_end_sample_stem(self) -> None:
        boat, archive_date = parse_workbook_filename("World_s_End-18-Sep-2026-06_38")
        self.assertEqual(boat, "World's End")
        self.assertEqual(archive_date, date(2026, 9, 18))

    def test_underscores_become_spaces(self) -> None:
        boat, archive_date = parse_workbook_filename("Sea_Breeze-01-Jan-2026-00_00")
        self.assertEqual(boat, "Sea Breeze")
        self.assertEqual(archive_date, date(2026, 1, 1))

    def test_hyphenated_boat_token(self) -> None:
        boat, _archive_date = parse_workbook_filename("My-Boat-18-Sep-2026-06_38")
        self.assertEqual(boat, "My-Boat")

    def test_unrecognized_name_returns_none(self) -> None:
        self.assertIsNone(parse_workbook_filename("outstanding"))

    def test_bad_month_raises(self) -> None:
        with self.assertRaises(ValueError):
            parse_workbook_filename("Boat-18-Xxx-2026-06_38")

    def test_bad_day_raises(self) -> None:
        with self.assertRaises(ValueError):
            parse_workbook_filename("Boat-32-Sep-2026-06_38")

    def test_override_when_name_has_no_date(self) -> None:
        boat, archive_date = resolve_archive_target("notes.xlsx", "Aurora", date(2026, 9, 27))
        self.assertEqual(boat, "Aurora")
        self.assertEqual(archive_date, date(2026, 9, 27))

    def test_pattern_wins_over_override(self) -> None:
        boat, archive_date = resolve_archive_target(
            "World_s_End-18-Sep-2026-06_38.xlsx",
            "Other",
            date(2026, 1, 1),
        )
        self.assertEqual(boat, "World's End")
        self.assertEqual(archive_date, date(2026, 9, 18))

    def test_missing_pattern_and_override_raises(self) -> None:
        with self.assertRaises(ValueError):
            resolve_archive_target("notes.xlsx", None, date(2026, 9, 27))


class DiscoverTests(unittest.TestCase):
    def test_one_workbook_and_any_pdf_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "Whatever-Boat.xlsx"
            workbook.write_bytes(b"")
            (root / "scan.PDF").write_bytes(b"")
            (root / "ΤΠΥ524.pdf").write_bytes(b"")
            (root / ".DS_Store").write_bytes(b"")
            (root / "~$Whatever-Boat.xlsx").write_bytes(b"")
            (root / "notes.txt").write_text("x", encoding="utf-8")
            found, pdfs = discover_inputs(root)
            self.assertEqual(found, workbook)
            self.assertEqual([path.name for path in pdfs], ["scan.PDF", "ΤΠΥ524.pdf"])

    def test_zero_or_two_workbooks_name_the_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with self.assertRaises(ValueError) as none_found:
                discover_inputs(root)
            self.assertIn("none", str(none_found.exception))
            (root / "a.xlsx").write_bytes(b"")
            (root / "b.XLSX").write_bytes(b"")
            with self.assertRaises(ValueError) as both:
                discover_inputs(root)
            self.assertIn("a.xlsx", str(both.exception))
            self.assertIn("b.XLSX", str(both.exception))


class WorkbookTests(unittest.TestCase):
    def test_reads_by_header_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "boat.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "NOTES",
                    "SUPPLIER NAME",
                    "INVOICE NUMBER",
                    "DESCRIPTION",
                    "AMOUNT ORIGINAL",
                    "CURRENCY ORIGIN",
                    "CREATE DATE",
                    "INVOICE/QUESTION DATE",
                    "WORK/ORDER STATUS",
                    "BOAT CATEGORY",
                    "BOAT SUBCATEGORY",
                ]
            )
            sheet.append(
                [
                    "ignore",
                    "PALAMAIKI",
                    "32 & 266",
                    "SAILS",
                    161.2,
                    "EUR",
                    "16 Sep 2026",
                    datetime(2026, 9, 17),
                    "Completed",
                    "Maintenance",
                    "Rigging",
                ]
            )
            sheet.append([None] * 11)
            sheet.append(
                [
                    None,
                    "EMPORIO",
                    "12",
                    "anchor",
                    "1,177.75",
                    "EUR",
                    "08 Sep 2026",
                    "08 Sep 2026",
                    "Completed",
                    "Equipment",
                    "General",
                ]
            )
            workbook.save(path)
            workbook.close()
            payments = load_payments(path)
        self.assertEqual(len(payments), 2)
        self.assertEqual(payments[0].supplier_name, "PALAMAIKI")
        self.assertEqual(payments[0].invoice_number, "32 & 266")
        self.assertEqual(payments[0].amount_original, Decimal("161.20"))
        self.assertEqual(payments[0].create_date, date(2026, 9, 16))
        self.assertEqual(payments[0].invoice_date, date(2026, 9, 17))
        self.assertEqual(payments[1].amount_original, Decimal("1177.75"))

    def test_missing_column_and_amount(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["SUPPLIER NAME"])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError) as missing:
                load_payments(path)
            self.assertIn("AMOUNT ORIGINAL", str(missing.exception))

            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "SUPPLIER NAME",
                    "INVOICE NUMBER",
                    "DESCRIPTION",
                    "AMOUNT ORIGINAL",
                    "CURRENCY ORIGIN",
                    "CREATE DATE",
                    "INVOICE/QUESTION DATE",
                    "WORK/ORDER STATUS",
                    "BOAT CATEGORY",
                    "BOAT SUBCATEGORY",
                ]
            )
            sheet.append(["Acme", "1", "job", None, "EUR", None, None, "New", "Ops", "Fuel"])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError) as blank_amount:
                load_payments(path)
            self.assertIn("Row 2", str(blank_amount.exception))

    def test_amount_formats(self) -> None:
        self.assertEqual(parse_amount("704,22"), Decimal("704.22"))
        self.assertEqual(parse_amount("1.177,75"), Decimal("1177.75"))
        self.assertEqual(parse_amount(310), Decimal("310.00"))


class MatchTests(unittest.TestCase):
    def test_tolerance(self) -> None:
        payment = _payment(amount_original=Decimal("100.00"), invoice_number="1")
        close = match_payments([payment], [_invoice("a.pdf", "100.01", "9")])
        self.assertEqual(close[0].invoices[0].source_name, "a.pdf")
        far = match_payments([payment], [_invoice("b.pdf", "100.02", "1")])
        self.assertFalse(far[0].matched)

    def test_tie_break_by_invoice_number_then_filename(self) -> None:
        payment = _payment(amount_original=Decimal("100.00"), invoice_number="1187")
        invoices = [
            _invoice("other.pdf", "100.00", "9999"),
            _invoice("right.pdf", "100.00", "1187"),
        ]
        matched = match_payments([payment], invoices)
        self.assertEqual(matched[0].invoices[0].source_name, "right.pdf")

        by_name = match_payments(
            [payment],
            [
                _invoice("9999.pdf", "100.00", None),
                _invoice("invoice-1187.pdf", "100.00", None),
            ],
        )
        self.assertEqual(by_name[0].invoices[0].source_name, "invoice-1187.pdf")

    def test_ambiguous_single_totals_stay_unmatched(self) -> None:
        payment = _payment(amount_original=Decimal("100.00"), invoice_number="PROFORMA")
        matched = match_payments(
            [payment],
            [_invoice("a.pdf", "100.00", "1"), _invoice("b.pdf", "100.00", "2")],
        )
        self.assertFalse(matched[0].matched)

    def test_single_total_beats_a_sum(self) -> None:
        payment = _payment(amount_original=Decimal("30.00"), invoice_number="10/20")
        matched = match_payments(
            [payment],
            [
                _invoice("full.pdf", "30.00", "999"),
                _invoice("10.pdf", "10.00", "10"),
                _invoice("20.pdf", "20.00", "20"),
            ],
        )
        self.assertEqual([invoice.source_name for invoice in matched[0].invoices], ["full.pdf"])

    def test_four_pdf_sum(self) -> None:
        payment = _payment(amount_original=Decimal("2236.96"), invoice_number="524/525/585/586")
        invoices = [
            _invoice("524.pdf", "500.00", "524"),
            _invoice("525.pdf", "600.00", "525"),
            _invoice("585.pdf", "700.00", "586"),
            _invoice("586.pdf", "436.96", "585"),
            _invoice("unused.pdf", "50.00", "1"),
        ]
        matched = match_payments([payment], invoices)
        self.assertEqual(
            [invoice.source_name for invoice in matched[0].invoices],
            ["524.pdf", "525.pdf", "585.pdf", "586.pdf"],
        )

    def test_prefers_combination_whose_numbers_intersect(self) -> None:
        payment = _payment(amount_original=Decimal("30.00"), invoice_number="10/20")
        matched = match_payments(
            [payment],
            [
                _invoice("10.pdf", "10.00", "10"),
                _invoice("20.pdf", "20.00", "20"),
                _invoice("12.pdf", "12.00", "1"),
                _invoice("18.pdf", "18.00", "2"),
            ],
        )
        self.assertEqual(
            [invoice.source_name for invoice in matched[0].invoices],
            ["10.pdf", "20.pdf"],
        )

    def test_ambiguous_sums_stay_unmatched(self) -> None:
        payment = _payment(amount_original=Decimal("30.00"), invoice_number="AA & BB")
        matched = match_payments(
            [payment],
            [
                _invoice("10.pdf", "10.00", "1"),
                _invoice("20.pdf", "20.00", "2"),
                _invoice("14.pdf", "14.00", "3"),
                _invoice("16.pdf", "16.00", "4"),
            ],
        )
        self.assertFalse(matched[0].matched)

    def test_only_summing_combination_matches_without_tokens(self) -> None:
        payment = _payment(amount_original=Decimal("30.00"), invoice_number="AA & BB")
        matched = match_payments(
            [payment],
            [_invoice("10.pdf", "10.00", "1"), _invoice("20.pdf", "20.00", "2")],
        )
        self.assertEqual(
            [invoice.source_name for invoice in matched[0].invoices],
            ["10.pdf", "20.pdf"],
        )

    def test_pdf_is_not_reused(self) -> None:
        first = _payment(row_number=2, amount_original=Decimal("100.00"), invoice_number="1")
        second = _payment(row_number=3, amount_original=Decimal("100.00"), invoice_number="2")
        matched = match_payments(
            [first, second],
            [_invoice("only.pdf", "100.00", "1"), _invoice("other.pdf", "40.00", "2")],
        )
        self.assertEqual(matched[0].invoices[0].source_name, "only.pdf")
        self.assertFalse(matched[1].matched)

    def test_more_than_six_invoice_parts_are_not_subset_matched(self) -> None:
        payment = _payment(
            amount_original=Decimal("70.00"),
            invoice_number="1/2/3/4/5/6/7",
        )
        invoices = [_invoice(f"{index}.pdf", "10.00", str(index)) for index in range(1, 8)]
        matched = match_payments([payment], invoices)
        self.assertFalse(matched[0].matched)

    def test_leading_zeros_count_as_the_same_number(self) -> None:
        payment = _payment(amount_original=Decimal("10.00"), invoice_number="32")
        matched = match_payments(
            [payment],
            [_invoice("a.pdf", "10.00", "99"), _invoice("b.pdf", "10.00", "0032")],
        )
        self.assertEqual(matched[0].invoices[0].source_name, "b.pdf")


class CsvTests(unittest.TestCase):
    def test_matched_row_sums_and_filenames(self) -> None:
        payment = _payment(
            description="2 seabob service",
            supplier_name="DERAMAR",
            invoice_number="524/525",
            amount_original=Decimal("37.20"),
        )
        matched = MatchedPayment(
            payment,
            (
                _invoice("a.pdf", "12.40", "524", net="10.00"),
                _invoice("b.pdf", "24.80", "525", net="20.00", vat="4.80"),
            ),
        )
        fields = to_csv_fields("World's End", matched)
        self.assertEqual(fields["Tasks"], "2 seabob service")
        self.assertEqual(fields["Supplier"], "DERAMAR")
        self.assertEqual(fields["Start Date"], "2026-09-01")
        self.assertEqual(fields["Delivery Date"], "2026-09-02")
        self.assertEqual(fields["Owner's Approval"], "")
        self.assertEqual(fields["Status"], "Completed")
        self.assertEqual(fields["Estimated Cost"], "37.20")
        self.assertEqual(fields["Net Price"], "30.00")
        self.assertEqual(fields["VAT 24%"], "7.20")
        self.assertEqual(fields["Final Cost"], "37.20")
        self.assertEqual(
            fields["Notes"],
            "World's End | Maintenance / Electrician | 524/525",
        )
        self.assertEqual(fields["Invoice"], "a.pdf; b.pdf")
        self.assertEqual(tuple(fields), CSV_HEADERS)

    def test_unmatched_row(self) -> None:
        fields = to_csv_fields("Aurora", MatchedPayment(_payment(), ()))
        self.assertEqual(fields["Net Price"], "")
        self.assertEqual(fields["VAT 24%"], "")
        self.assertEqual(fields["Final Cost"], "100.00")
        self.assertEqual(fields["Invoice"], "")
        self.assertTrue(fields["Notes"].endswith("| UNMATCHED"))

    def test_net_blank_when_any_invoice_has_no_net(self) -> None:
        matched = MatchedPayment(
            _payment(amount_original=Decimal("30.00"), invoice_number="1/2"),
            (
                _invoice("a.pdf", "10.00", net="8.00", vat="2.00"),
                _invoice("b.pdf", "20.00"),
            ),
        )
        fields = to_csv_fields("Aurora", matched)
        self.assertEqual(fields["Net Price"], "")
        self.assertEqual(fields["VAT 24%"], "")

    def test_write_csv_round_trip(self) -> None:
        import csv

        payment = _payment(description="batteries")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.csv"
            write_csv(path, "Aurora", [MatchedPayment(payment, ())])
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
        self.assertEqual(list(rows[0]), list(CSV_HEADERS))
        self.assertEqual(rows[0]["Tasks"], "batteries")
        self.assertEqual(rows[0]["Invoice"], "")
        self.assertTrue(rows[0]["Notes"].endswith("| UNMATCHED"))


class ExtractionTests(unittest.TestCase):
    def test_to_extracted_invoice_rounds_money(self) -> None:
        extraction = InvoiceExtraction(
            supplier_name="  Acme  ",
            invoice_number=1187,
            invoice_date="2026-09-03",
            currency="EUR",
            net_amount="100,00",
            vat_amount=None,
            total_amount="124.00",
            description="radio",
        )
        invoice = to_extracted_invoice("scan.pdf", extraction)
        self.assertEqual(invoice.supplier_name, "Acme")
        self.assertEqual(invoice.invoice_number, "1187")
        self.assertEqual(invoice.net_amount, Decimal("100.00"))
        self.assertIsNone(invoice.vat_amount)
        self.assertEqual(invoice.total_amount, Decimal("124.00"))
        self.assertEqual(invoice.source_name, "scan.pdf")

    def test_retry_after_parses_free_tier_message(self) -> None:
        message = (
            "429 RESOURCE_EXHAUSTED. Quota exceeded for metric: "
            "generativelanguage.googleapis.com/generate_content_free_tier_requests, "
            "limit: 5, model: gemini-3.8-flash. Please retry in 44.681195921s."
        )
        self.assertAlmostEqual(retry_after_seconds(RuntimeError(message)) or 0, 44.681195921)


if __name__ == "__main__":
    unittest.main()
